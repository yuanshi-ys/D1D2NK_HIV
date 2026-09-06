import os
import numpy as np
import pandas as pd
import openpyxl
from itertools import combinations
from scipy.stats import entropy, fisher_exact

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.path.dirname(SCRIPT_DIR)
PERIS_PATH = os.path.join(UPLOAD_DIR, "output", "perIS_lite.csv")
INTERSECT_PATH = os.path.join(UPLOAD_DIR, "output", "perIS_intersect.csv")
DISTANCE_PATH = os.path.join(UPLOAD_DIR, "output", "perIS_distance.csv")
OUT_DIR = os.path.join(UPLOAD_DIR, "source_data")

TREATMENT_ORDER = ["No NK", "GFP-NK", "D1D2-NK"]
ACTIVE_COLS = ["In_Gene", "In_TF_binding_site", "In_enhancer", "In_promoter"]
HISTONE_MARKS = ["H3K36me3", "H3K27ac", "H3K4me1", "H3K9me3", "H3K27me3", "H3K4me3"]

TITLES = {
    "Fig. 5a": "Number of proviral barcodes associated with rebound viremia",
    "Fig. 5b": "Genetic diversity of proviral barcodes associated with rebound viremia ",
    "Fig. 5c": "Dominance of proviral barcodes associated with rebound viremia",
    "Fig. 5d": "Fraction of shared proviral BC, IS-UMI weighted",
    "Fig. 5e": "UMI per IS lineage, proliferated (expanded, UMI>1) clones",
    "Fig. 5f": "UMI per IS lineage, proliferated clones associated with rebound viremia",
    "Fig. 5g": "Fraction of shared IS lineage, Jaccard",
    "Fig. 5h": "Fraction of shared IS lineage, Jaccard, by organ pair",
    "Fig. 6a": "Mosaic plots of proviruses associated with rebound viremia integrated within genes -- collapsed to unique lineages (organ collapsed; same site with different barcode still counted as independent integration event)",
    "Fig. 6b": "Mosaic plots of proviruses associated with rebound viremia integrated within transcriptionally active regions -- collapsed to unique lineages (organ collapsed; same site with different barcode still counted as independent integration event)",
    "Fig. 6c": "Distribution of rebound-associated proviruses in activating or repressive histone regions -- collapsed to unique lineages (organ collapsed; barcode kept)",
    "Fig. 6d": "Distance of rebound-associated proviruses to the nearest histone mark of GFP-NK vs No NK group -- collapsed to unique lineages",
    "Fig. 6e": "Distance of rebound-associated proviruses to the nearest histone mark of D1D2-NK vs No NK group -- collapsed to unique lineages",
    "Fig. S7a": "The number of all proviral barcodes were quantified by deep sequencing.",
    "Fig. S7b": "The genetic diversity of all proviral barcodes were quantified by deep sequencing. ",
    "Fig. S7c": "The dominance of all proviral barcodes were quantified by deep sequencing",
    "Fig. S7d": "Fraction of shared proviral BC, IS-UMI weighted -- all barcodes (not restricted to rebound viremia)",
    "Fig. S7e": "Fraction of shared proviral BC, IS-UMI weighted, associated with rebound viremia -- by organ pair",
    "Fig. S8a": "UMI per IS, tissue-level proliferated clones (Proliferation_tissue, per organ-sample, not lineage-summed)",
    "Fig. S8b": "UMI per IS, tissue-level proliferated clones associated with rebound viremia (per organ-sample, not lineage-summed)",
    "Fig. S8c": "Fraction of shared IS lineage, Jaccard, proliferated (expanded, UMI>1) clones",
}


def load_perIS(path):
    perIS = pd.read_csv(path)
    if "SampleID" not in perIS.columns:
        perIS["SampleID"] = perIS["AnimalKey"] + "-" + perIS["Organ"]
    return perIS


def get_shannon(umi_by_barcode):
    freq = umi_by_barcode / np.sum(umi_by_barcode)
    return entropy(freq, base=2)


def get_dominance(umi_by_barcode):
    arr = np.asarray(umi_by_barcode, dtype=float)
    t = arr.sum()
    if t == 0:
        return np.nan
    return np.sum((arr / t) ** 2)


def _per_sample_umi_by_barcode(perIS):
    return perIS.groupby(["SampleID", "barcode"])["UMI"].sum().reset_index()


def fig5a_num_proviral_bc(perIS, viremia_only):
    all_samples = perIS.drop_duplicates("SampleID")[["SampleID", "Treatment"]]
    df = perIS[perIS["Viremia"]] if viremia_only else perIS
    counts = df.groupby("SampleID")["barcode"].nunique()
    per_sample = all_samples.copy()
    per_sample["n_proviral_barcodes"] = per_sample["SampleID"].map(counts).fillna(0).astype(int)
    return per_sample


def fig5b_shannon(perIS, viremia_only):
    all_samples = perIS.drop_duplicates("SampleID")[["SampleID", "Treatment"]]
    df = perIS[perIS["Viremia"]] if viremia_only else perIS
    bc_umi = _per_sample_umi_by_barcode(df)
    shannon_by_sample = bc_umi.groupby("SampleID")["UMI"].apply(lambda s: get_shannon(s.values))
    out = all_samples.copy()
    out["Shannon"] = out["SampleID"].map(shannon_by_sample)
    return out.dropna(subset=["Shannon"]).reset_index(drop=True)


def fig5c_dominance(perIS, viremia_only):
    all_samples = perIS.drop_duplicates("SampleID")[["SampleID", "Treatment"]]
    df = perIS[perIS["Viremia"]] if viremia_only else perIS
    bc_umi = _per_sample_umi_by_barcode(df)
    dominance_by_sample = bc_umi.groupby("SampleID")["UMI"].apply(lambda s: get_dominance(s.values))
    out = all_samples.copy()
    out["Dominance"] = out["SampleID"].map(dominance_by_sample)
    return out.dropna(subset=["Dominance"]).reset_index(drop=True)


def _find_data(animal, organ, df):
    return df[(df["AnimalKey"] == animal) & (df["Organ"] == organ)]


def _get_jaccard_index_is_bc(data1, data2):
    key1 = set(zip(data1["site"], data1["barcode"]))
    key2 = set(zip(data2["site"], data2["barcode"]))
    overlap = key1 & key2
    denom = data1["UMI"].sum()
    if denom == 0:
        return float("nan")
    keys1 = list(zip(data1["site"], data1["barcode"]))
    mask = [k in overlap for k in keys1]
    return data1.loc[mask, "UMI"].sum() / denom


def fig5d_isbc_umi_weighted(perIS, viremia_only):
    df = perIS[perIS["Viremia"]] if viremia_only else perIS
    summary = df.drop_duplicates("AnimalKey")[["AnimalKey", "Treatment"]]
    rows = []
    for animal in df["AnimalKey"].unique():
        organs = df[df["AnimalKey"] == animal]["Organ"].unique()
        for donor_organ in organs:
            for acceptor_organ in organs:
                if donor_organ == acceptor_organ:
                    continue
                data1 = _find_data(animal, donor_organ, df)
                data2 = _find_data(animal, acceptor_organ, df)
                rows.append({
                    "organ1": donor_organ, "organ2": acceptor_organ, "AnimalKey": animal,
                    "value": _get_jaccard_index_is_bc(data1, data2),
                })
    jaccard = pd.DataFrame(rows)
    if jaccard.empty:
        return jaccard
    return pd.merge(summary, jaccard, on="AnimalKey")[
        ["Treatment", "AnimalKey", "organ1", "organ2", "value"]
    ].drop_duplicates()


def _jaccard_set_bc_is(data1, data2):
    key1 = set(zip(data1["barcode"], data1["site"]))
    key2 = set(zip(data2["barcode"], data2["site"]))
    union = key1 | key2
    if not union:
        return float("nan")
    return len(key1 & key2) / len(union)


def fig_setjaccard_by_animal(perIS):
    summary = perIS.drop_duplicates("AnimalKey")[["AnimalKey", "Treatment"]]
    rows = []
    for animal in perIS["AnimalKey"].unique():
        organs = perIS[perIS["AnimalKey"] == animal]["Organ"].unique()
        for organ1, organ2 in combinations(sorted(organs), 2):
            data1 = _find_data(animal, organ1, perIS)
            data2 = _find_data(animal, organ2, perIS)
            rows.append({
                "organ1": organ1, "organ2": organ2, "AnimalKey": animal,
                "value": _jaccard_set_bc_is(data1, data2),
            })
    jac = pd.DataFrame(rows)
    if jac.empty:
        return jac
    return pd.merge(summary, jac, on="AnimalKey")[
        ["Treatment", "AnimalKey", "organ1", "organ2", "value"]
    ].drop_duplicates()


def _add_organpair_label(g):
    order = ["Sp", "BM", "Liv"]
    rename = {"SP": "Sp", "LIV": "Liv", "BM": "BM"}

    def pair_label(o1, o2):
        a, b = sorted([rename.get(o1, o1), rename.get(o2, o2)], key=lambda x: order.index(x))
        return f"{a}-{b}"

    g = g.copy()
    g["pair"] = g.apply(lambda r: pair_label(r["organ1"], r["organ2"]), axis=1)
    return g


def organpair_wide(df_with_pair, value_col):
    pairs = ["Sp-Liv", "Sp-BM", "BM-Liv"]
    cols = {}
    for pair in pairs:
        for t in TREATMENT_ORDER:
            t_label = {"No NK": "No NK", "GFP-NK": "GFP", "D1D2-NK": "D1D2"}[t]
            key = f"{pair} ({t_label})"
            sub = df_with_pair[(df_with_pair["pair"] == pair) & (df_with_pair["Treatment"] == t)][value_col]
            cols[key] = sub.reset_index(drop=True)
    return pd.DataFrame(cols)


def perIS_by_animal_lineage(perIS):
    return perIS.groupby(["AnimalKey", "site", "barcode"], as_index=False).agg(
        UMI=("UMI", "sum"), Treatment=("Treatment", "first"), Viremia=("Viremia", "first"))


def lineage_umi(perIS, viremia_only):
    lineage = perIS_by_animal_lineage(perIS)
    expanded = lineage[lineage["UMI"] > 1]
    if viremia_only:
        expanded = expanded[expanded["Viremia"]]
    return expanded[["Treatment", "AnimalKey", "site", "barcode", "UMI"]].copy()


def tissue_prolif_umi(perIS, viremia_only):
    df = perIS[perIS["Proliferation_tissue"]]
    if viremia_only:
        df = df[df["Viremia"]]
    return df[["Treatment", "AnimalKey", "site", "barcode", "SampleID", "UMI"]].copy()


def wide_ragged(df, value_col, group_col="Treatment"):
    cols = {}
    for t in TREATMENT_ORDER:
        cols[t] = df[df[group_col] == t][value_col].reset_index(drop=True)
    return pd.DataFrame(cols)


def write_block(ws, title, header, data_rows, start_row=1):
    ws.cell(row=start_row, column=1, value=title)
    header_row = start_row + 2
    for j, h in enumerate(header, start=1):
        ws.cell(row=header_row, column=j, value=h)
    for i, row in enumerate(data_rows):
        for j, v in enumerate(row, start=1):
            if pd.notna(v):
                ws.cell(row=header_row + 1 + i, column=j, value=v)


def write_wide_sheet(wb, sheet_name, title, wide_df):
    ws = wb.create_sheet(sheet_name)
    write_block(ws, title, list(wide_df.columns), wide_df.values.tolist())


def write_mosaic_sheet(wb, sheet_name, title, vir, col, label_in, label_out):
    ws = wb.create_sheet(sheet_name)
    ws.cell(row=1, column=1, value=title)
    row = 3
    for t in TREATMENT_ORDER:
        sub = vir[vir["Treatment"] == t][col]
        n_total = len(sub)
        n_in = int(sub.sum())
        frac_in = sub.mean()
        ws.cell(row=row, column=1, value=f"{t} {label_in}")
        ws.cell(row=row, column=2, value=f"{t} {label_out}")
        ws.cell(row=row + 1, column=1, value=frac_in)
        ws.cell(row=row + 1, column=2, value=1 - frac_in)
        ws.cell(row=row + 2, column=1, value="n_total")
        ws.cell(row=row + 2, column=2, value=n_total)
        ws.cell(row=row + 3, column=1, value="n_in")
        ws.cell(row=row + 3, column=2, value=n_in)
        row += 5


def write_fig6c_sheet(wb, title, odds_df):
    ws = wb.create_sheet("Fig. 6c")
    ws.cell(row=1, column=1, value=title)
    ws.cell(row=2, column=1, value="Odd ratio")
    ws.cell(row=2, column=2, value="D1D2-NK vs No NK")
    ws.cell(row=2, column=3, value="GFP-NK vs No NK")
    for i, mark in enumerate(HISTONE_MARKS):
        ws.cell(row=3 + i, column=1, value=mark)
        ws.cell(row=3 + i, column=2, value=odds_df.loc[mark, "OR_D1D2_vs_NoNK"])
        ws.cell(row=3 + i, column=3, value=odds_df.loc[mark, "OR_GFPNK_vs_NoNK"])
    prow = 3 + len(HISTONE_MARKS) + 1
    ws.cell(row=prow, column=1, value="p values")
    ws.cell(row=prow, column=2, value="D1D2-NK vs No NK")
    ws.cell(row=prow, column=3, value="GFP-NK vs No NK")
    for i, mark in enumerate(HISTONE_MARKS):
        ws.cell(row=prow + 1 + i, column=1, value=mark)
        ws.cell(row=prow + 1 + i, column=2, value=odds_df.loc[mark, "p_D1D2_vs_NoNK"])
        ws.cell(row=prow + 1 + i, column=3, value=odds_df.loc[mark, "p_GFPNK_vs_NoNK"])


def write_fig6de_sheet(wb, sheet_name, title, vir, group, distance_fields):
    ws = wb.create_sheet(sheet_name)
    ws.cell(row=1, column=1, value=title)
    ws.cell(row=2, column=1, value="Treatment")
    ws.cell(row=2, column=2, value="Histone_mark")
    ws.cell(row=2, column=3, value="Distance to nearest histone (Log10 bp)")
    i = 0
    for mark in distance_fields:
        field = f"distance_{mark}"
        for t in [group, "No NK"]:
            sub = vir[vir["Treatment"] == t][field].dropna()
            sub = sub[sub != 0]
            for val in sub:
                ws.cell(row=3 + i, column=1, value=t)
                ws.cell(row=3 + i, column=2, value=mark)
                ws.cell(row=3 + i, column=3, value=float(np.log10(abs(val))))
                i += 1


def add_active_col(df):
    df = df.copy()
    df["In_Active"] = df[ACTIVE_COLS].any(axis=1)
    return df


def add_histone_in_cols(df, cutoff=2000):
    df = df.copy()
    for mark in HISTONE_MARKS:
        col = f"distance_{mark}"
        df[f"In_{mark}"] = df[col].notna() & (df[col].abs() <= cutoff)
    return df


def collapse_to_unique_lineage(df, check_cols):
    g = df.groupby(["AnimalID", "site", "barcode"])[check_cols].nunique()
    for c in check_cols:
        assert (g[c] > 1).sum() == 0
    return df.drop_duplicates(["AnimalID", "site", "barcode"]).copy()


def fisher_vs_baseline(df, col, group, baseline="No NK"):
    a = df[df["Treatment"] == group]
    b = df[df["Treatment"] == baseline]
    table = [[int(a[col].sum()), int((~a[col]).sum())], [int(b[col].sum()), int((~b[col]).sum())]]
    orr, p = fisher_exact(table, alternative="two-sided")
    return orr, p, len(a), len(b)


def odds_ratio_table(df, cols_and_labels):
    rows = []
    for col, label in cols_and_labels:
        or1, p1, n1, n0 = fisher_vs_baseline(df, col, "D1D2-NK")
        or2, p2, n2, _ = fisher_vs_baseline(df, col, "GFP-NK")
        rows.append({
            "Element": label,
            "frac_NoNK": df[df["Treatment"] == "No NK"][col].mean(),
            "frac_GFPNK": df[df["Treatment"] == "GFP-NK"][col].mean(),
            "frac_D1D2NK": df[df["Treatment"] == "D1D2-NK"][col].mean(),
            "OR_D1D2_vs_NoNK": or1, "p_D1D2_vs_NoNK": p1,
            "OR_GFPNK_vs_NoNK": or2, "p_GFPNK_vs_NoNK": p2,
            "n_NoNK": n0, "n_GFPNK": n2, "n_D1D2NK": n1,
        })
    return pd.DataFrame(rows).set_index("Element")


def make_fig5_workbook(perIS, out_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    a = fig5a_num_proviral_bc(perIS, True)
    write_wide_sheet(wb, "Fig. 5a", TITLES["Fig. 5a"], wide_ragged(a, "n_proviral_barcodes"))

    b = fig5b_shannon(perIS, True)
    write_wide_sheet(wb, "Fig. 5b", TITLES["Fig. 5b"], wide_ragged(b, "Shannon"))

    c = fig5c_dominance(perIS, True)
    write_wide_sheet(wb, "Fig. 5c", TITLES["Fig. 5c"], wide_ragged(c, "Dominance"))

    d = fig5d_isbc_umi_weighted(perIS, True)
    write_wide_sheet(wb, "Fig. 5d", TITLES["Fig. 5d"], wide_ragged(d, "value"))

    e = lineage_umi(perIS, False)
    write_wide_sheet(wb, "Fig. 5e", TITLES["Fig. 5e"], wide_ragged(e, "UMI"))

    f = lineage_umi(perIS, True)
    write_wide_sheet(wb, "Fig. 5f", TITLES["Fig. 5f"], wide_ragged(f, "UMI"))

    g = fig_setjaccard_by_animal(perIS)
    write_wide_sheet(wb, "Fig. 5g", TITLES["Fig. 5g"], wide_ragged(g, "value"))

    h = _add_organpair_label(g)
    write_wide_sheet(wb, "Fig. 5h", TITLES["Fig. 5h"], organpair_wide(h, "value"))

    wb.save(out_path)


def make_fig6_workbook(intersect_path, distance_path, out_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    vir_i = add_active_col(pd.read_csv(intersect_path))
    vir_i = vir_i[vir_i["Viremia"]].copy()
    vir_i = collapse_to_unique_lineage(vir_i, ["In_Gene", "In_Active"])
    write_mosaic_sheet(wb, "Fig. 6a", TITLES["Fig. 6a"], vir_i, "In_Gene", "in genes", "Not in genes")
    write_mosaic_sheet(wb, "Fig. 6b", TITLES["Fig. 6b"], vir_i, "In_Active", "active", "inactive")

    vir_d = pd.read_csv(distance_path)
    vir_d = vir_d[vir_d["Viremia"]].copy()
    vir_d = collapse_to_unique_lineage(vir_d, [f"distance_{m}" for m in HISTONE_MARKS])
    vir_d = add_histone_in_cols(vir_d)
    odds = odds_ratio_table(vir_d, [(f"In_{m}", m) for m in HISTONE_MARKS])
    write_fig6c_sheet(wb, TITLES["Fig. 6c"], odds)

    write_fig6de_sheet(wb, "Fig. 6d", TITLES["Fig. 6d"], vir_d, "GFP-NK", HISTONE_MARKS)
    write_fig6de_sheet(wb, "Fig. 6e", TITLES["Fig. 6e"], vir_d, "D1D2-NK", HISTONE_MARKS)

    wb.save(out_path)


def make_figS7_workbook(perIS, out_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    a = fig5a_num_proviral_bc(perIS, False)
    write_wide_sheet(wb, "Fig. S7a", TITLES["Fig. S7a"], wide_ragged(a, "n_proviral_barcodes"))

    b = fig5b_shannon(perIS, False)
    write_wide_sheet(wb, "Fig. S7b", TITLES["Fig. S7b"], wide_ragged(b, "Shannon"))

    c = fig5c_dominance(perIS, False)
    write_wide_sheet(wb, "Fig. S7c", TITLES["Fig. S7c"], wide_ragged(c, "Dominance"))

    d = fig5d_isbc_umi_weighted(perIS, False)
    write_wide_sheet(wb, "Fig. S7d", TITLES["Fig. S7d"], wide_ragged(d, "value"))

    e = fig5d_isbc_umi_weighted(perIS, True)
    e_pair = _add_organpair_label(e)
    write_wide_sheet(wb, "Fig. S7e", TITLES["Fig. S7e"], organpair_wide(e_pair, "value"))

    wb.save(out_path)


def make_figS8_workbook(perIS, out_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    a = tissue_prolif_umi(perIS, False)
    write_wide_sheet(wb, "Fig. S8a", TITLES["Fig. S8a"], wide_ragged(a, "UMI"))

    b = tissue_prolif_umi(perIS, True)
    write_wide_sheet(wb, "Fig. S8b", TITLES["Fig. S8b"], wide_ragged(b, "UMI"))

    exp_only = perIS[perIS["UMI"] > 1]
    c = fig_setjaccard_by_animal(exp_only)
    write_wide_sheet(wb, "Fig. S8c", TITLES["Fig. S8c"], wide_ragged(c, "value"))

    wb.save(out_path)


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    perIS = load_perIS(PERIS_PATH)
    make_fig5_workbook(perIS, os.path.join(OUT_DIR, "Source Data Fig. 5.xlsx"))
    make_fig6_workbook(INTERSECT_PATH, DISTANCE_PATH, os.path.join(OUT_DIR, "Source Data Fig. 6.xlsx"))
    make_figS7_workbook(perIS, os.path.join(OUT_DIR, "Source Data Fig. S7.xlsx"))
    make_figS8_workbook(perIS, os.path.join(OUT_DIR, "Source Data Fig. S8.xlsx"))


if __name__ == "__main__":
    main()
